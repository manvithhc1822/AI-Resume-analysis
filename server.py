from flask import Flask, request, jsonify, send_from_directory
from werkzeug.utils import secure_filename
from pathlib import Path
import re

from PyPDF2 import PdfReader
from docx import Document

app = Flask(__name__)

UPLOAD_FOLDER = Path("uploads")
UPLOAD_FOLDER.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}

SKILLS = [
    "python", "java", "javascript", "typescript", "html", "css",
    "react", "angular", "vue", "bootstrap",
    "flask", "django", "fastapi", "node.js", "node",
    "sql", "mysql", "postgresql", "mongodb", "oracle",
    "git", "github", "docker", "linux", "azure", "aws",
    "rest api", "api", "json", "postman",
    "pandas", "numpy", "matplotlib", "power bi", "excel",
    "machine learning", "artificial intelligence", "opencv",
    "testing", "selenium", "oops", "data structures",
    "php", "laravel", "android", "kotlin", "firebase"
]

JOB_ROLES = {
    "Python Developer": [
        "python", "flask", "django", "sql", "git", "rest api"
    ],
    "Frontend Developer": [
        "html", "css", "javascript", "react", "bootstrap", "git"
    ],
    "Backend Developer": [
        "python", "java", "sql", "rest api", "git", "database"
    ],
    "Full Stack Developer": [
        "html", "css", "javascript", "python", "flask", "sql", "git"
    ],
    "Data Analyst": [
        "python", "sql", "excel", "pandas", "numpy", "matplotlib", "power bi"
    ],
    "Software Developer": [
        "python", "java", "sql", "git", "oops", "data structures"
    ],
    "QA / Test Engineer": [
        "testing", "selenium", "python", "java", "sql", "git"
    ]
}


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def extract_text(file_path):
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(str(file_path))
        text = []

        for page in reader.pages:
            page_text = page.extract_text()

            if page_text:
                text.append(page_text)

        return "\n".join(text)

    if extension == ".docx":
        document = Document(str(file_path))
        text = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text.append(cell.text)

        return "\n".join(text)

    if extension == ".txt":
        return file_path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

    return ""


def contains_skill(text, skill):
    text = text.lower()

    if skill in ["c", "api"]:
        pattern = r"\b" + re.escape(skill) + r"\b"
        return bool(re.search(pattern, text))

    return skill.lower() in text


def detect_skills(text):
    found = []

    for skill in SKILLS:
        if contains_skill(text, skill):
            found.append(skill)

    return found


def find_email(text):
    pattern = r"[\w.+-]+@[\w-]+\.[\w.-]+"
    match = re.search(pattern, text)

    if match:
        return match.group(0)

    return ""


def find_phone(text):
    pattern = r"(?:\+91[\s-]?)?[6-9]\d{9}"
    match = re.search(pattern, text)

    if match:
        return match.group(0)

    return ""


def calculate_score(text, skills):
    lower_text = text.lower()

    score = 0

    if len(text) >= 500:
        score += 15
    elif len(text) >= 250:
        score += 8

    sections = [
        "education",
        "experience",
        "skills",
        "projects",
        "project",
        "certification",
        "certificate",
        "contact"
    ]

    for section in sections:
        if section in lower_text:
            score += 5

    score += min(len(skills) * 3, 30)

    if find_email(text):
        score += 5

    if find_phone(text):
        score += 5

    action_words = [
        "developed",
        "created",
        "implemented",
        "designed",
        "built",
        "managed",
        "optimized"
    ]

    action_count = 0

    for word in action_words:
        if word in lower_text:
            action_count += 1

    score += min(action_count * 2, 10)

    return min(score, 100)


def calculate_job_matches(skills):
    results = []

    skill_set = set(skills)

    for role, required_skills in JOB_ROLES.items():

        matched = []

        for skill in required_skills:
            if skill in skill_set:
                matched.append(skill)

        missing = []

        for skill in required_skills:
            if skill not in skill_set:
                missing.append(skill)

        percentage = int(
            (len(matched) / len(required_skills)) * 100
        )

        results.append({
            "role": role,
            "match": percentage,
            "matched": matched,
            "missing": missing
        })

    results.sort(
        key=lambda item: item["match"],
        reverse=True
    )

    return results


def generate_suggestions(text, skills):
    lower_text = text.lower()
    suggestions = []

    if not find_email(text):
        suggestions.append(
            "Add a professional email address."
        )

    if not find_phone(text):
        suggestions.append(
            "Add your phone number."
        )

    if "github" not in lower_text:
        suggestions.append(
            "Add your GitHub profile."
        )

    if "linkedin" not in lower_text:
        suggestions.append(
            "Add your LinkedIn profile."
        )

    if "project" not in lower_text:
        suggestions.append(
            "Add relevant projects with technologies used."
        )

    if "education" not in lower_text:
        suggestions.append(
            "Add your education details."
        )

    if "experience" not in lower_text and "internship" not in lower_text:
        suggestions.append(
            "Add internship, training, or practical experience."
        )

    if len(skills) < 5:
        suggestions.append(
            "Add relevant technical skills that you actually know."
        )

    if not any(
        word in lower_text
        for word in [
            "developed",
            "created",
            "implemented",
            "designed",
            "built"
        ]
    ):
        suggestions.append(
            "Use action words such as developed, built and implemented."
        )

    if not suggestions:
        suggestions.append(
            "Your resume contains the main sections. Keep achievements clear and measurable."
        )

    return suggestions[:8]


@app.route("/")
def home():
    return send_from_directory(".", "index.html")


@app.route("/style.css")
def css():
    return send_from_directory(".", "style.css")


@app.route("/script.js")
def javascript():
    return send_from_directory(".", "script.js")


@app.route("/analyze", methods=["POST"])
def analyze_resume():

    if "resume" not in request.files:
        return jsonify({
            "error": "Please select a resume."
        }), 400

    file = request.files["resume"]

    if file.filename == "":
        return jsonify({
            "error": "Please select a resume."
        }), 400

    if not allowed_file(file.filename):
        return jsonify({
            "error": "Only PDF, DOCX and TXT files are allowed."
        }), 400

    filename = secure_filename(file.filename)

    file_path = UPLOAD_FOLDER / filename

    try:
        file.save(file_path)

        text = extract_text(file_path).strip()

        if not text:
            return jsonify({
                "error": "No readable text found in the resume."
            }), 400

        skills = detect_skills(text)

        score = calculate_score(
            text,
            skills
        )

        job_matches = calculate_job_matches(
            skills
        )

        suggestions = generate_suggestions(
            text,
            skills
        )

        words = re.findall(
            r"\b[\w+#.-]+\b",
            text
        )

        keywords = []

        for word in words:
            clean_word = word.lower()

            if len(clean_word) >= 5:
                if clean_word not in keywords:
                    keywords.append(clean_word)

        keywords = keywords[:40]

        result = {
            "filename": filename,
            "score": score,
            "skills": skills,
            "email": find_email(text),
            "phone": find_phone(text),
            "word_count": len(words),
            "keywords": keywords,
            "job_matches": job_matches[:5],
            "suggestions": suggestions
        }

        return jsonify(result)

    except Exception as error:

        return jsonify({
            "error": str(error)
        }), 500

    finally:

        if file_path.exists():
            file_path.unlink()


if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )